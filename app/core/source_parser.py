import io
import uuid

from openpyxl import load_workbook

from app.core.session_store import (
    SourceFileMeta,
    SourceGroupMeta,
    put_source_bytes,
)
from app.core.template_parser import count_leading_blank_rows


def _parse_headers(file_bytes: bytes) -> tuple[int, dict[str, int], list[str]]:
    """
    Parse source file headers using read_only mode (memory-efficient for large files).
    Returns (header_row_1based, col_name->0based_row_tuple_idx, ordered_col_names).
    """
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    ws = wb.active
    if ws is None:
        wb.close()
        raise ValueError("Source file has no active sheet")

    blank_rows = count_leading_blank_rows(ws)
    header_row = blank_rows + 1

    col_names: list[str] = []
    col_to_idx: dict[str, int] = {}

    for row in ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True):
        for idx, val in enumerate(row):
            if val is not None:
                name = str(val)
                if name not in col_to_idx:  # keep first occurrence on duplicate headers
                    col_to_idx[name] = idx
                    col_names.append(name)

    wb.close()

    if not col_names:
        raise ValueError(f"Source file has no column headers (checked row {header_row})")

    return header_row, col_to_idx, col_names


def parse_source_files(files: list[tuple[bytes, str]]) -> list[SourceGroupMeta]:
    """
    Parse and group source files by column header structure (frozenset equality).
    Files with the same set of column names go into one group.
    Stores raw bytes in session_store for later conversion.
    Returns list of SourceGroupMeta (one per unique header set).
    """
    groups: dict[frozenset, SourceGroupMeta] = {}

    for file_bytes, filename in files:
        header_row, col_to_idx, col_names = _parse_headers(file_bytes)

        bytes_key = str(uuid.uuid4())
        put_source_bytes(bytes_key, file_bytes)

        file_meta = SourceFileMeta(
            file_id=str(uuid.uuid4()),
            filename=filename,
            source_columns=col_names,
            col_name_to_row_index=col_to_idx,
            header_row=header_row,
            bytes_key=bytes_key,
        )

        key = frozenset(col_names)
        if key not in groups:
            groups[key] = SourceGroupMeta(
                group_id=str(uuid.uuid4()),
                files=[file_meta],
                source_columns=col_names,  # first file's ordering is canonical
            )
        else:
            groups[key].files.append(file_meta)

    return list(groups.values())


def auto_map(template_columns: list[str], source_columns: list[str]) -> list[dict]:
    """Auto-map template columns to source columns by exact name match."""
    source_set = set(source_columns)
    return [
        {
            "output_col": col,
            "source_col": col if col in source_set else None,
            "is_extra": False,
        }
        for col in template_columns
    ]
