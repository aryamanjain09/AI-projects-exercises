import io

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.core.session_store import TemplateMetadata
from app.core.template_parser import count_leading_blank_rows, count_leading_blank_cols


def _parse_source(source_bytes: bytes) -> tuple[int, int, dict[str, int]]:
    """
    Parse a source .xlsx file and return:
      (header_row, leading_blank_cols, col_name_to_1based_col_index)
    Uses the same blank-detection logic as the template parser.
    """
    wb = load_workbook(io.BytesIO(source_bytes), data_only=True)
    ws = wb.active

    if ws is None or not ws.max_row or not ws.max_column:
        raise ValueError("Source sheet has no data")

    blank_rows = count_leading_blank_rows(ws)
    header_row = blank_rows + 1

    if header_row > ws.max_row:
        raise ValueError("Source sheet has no data")

    blank_cols = count_leading_blank_cols(ws, header_row)

    col_map: dict[str, int] = {}
    col = blank_cols + 1
    while col <= ws.max_column:
        val = ws.cell(row=header_row, column=col).value
        if val is None:
            break
        name = str(val)
        if name not in col_map:  # keep first occurrence on duplicate source headers
            col_map[name] = col
        col += 1

    return header_row, blank_cols, col_map


def _clear_data_rows(ws: Worksheet, header_row: int) -> None:
    """Set all cell values below the header row to None (preserves styles)."""
    if ws.max_row is None or ws.max_column is None:
        return
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.value = None


def convert(source_bytes: bytes, meta: TemplateMetadata) -> tuple[bytes, list[str]]:
    """
    Convert a source .xlsx to match the template structure.
    Returns (output_xlsx_bytes, warnings).
    """
    # Load a fresh copy of the template to preserve all styles
    out_wb = load_workbook(io.BytesIO(meta.template_wb_bytes))
    out_ws = out_wb.active

    # Load source (values only for reading)
    src_wb = load_workbook(io.BytesIO(source_bytes), data_only=True)
    src_ws = src_wb.active

    if out_ws is None:
        raise ValueError("Template workbook has no active sheet")
    if src_ws is None:
        raise ValueError("Source workbook has no active sheet")

    src_header_row, _src_blank_cols, src_col_map = _parse_source(source_bytes)

    # Clear existing data rows in the output (template may have sample data)
    _clear_data_rows(out_ws, meta.header_row)

    template_columns = meta.column_names
    template_col_set = set(template_columns)
    source_col_set = set(src_col_map.keys())

    warnings: list[str] = []
    for name in source_col_set - template_col_set:
        warnings.append(f"Source column '{name}' not in template — dropped")
    for name in template_col_set - source_col_set:
        warnings.append(f"Template column '{name}' not found in source — left blank")

    # Write data rows
    out_row = meta.header_row + 1
    src_max_row = src_ws.max_row or 0

    for src_row_idx in range(src_header_row + 1, src_max_row + 1):
        # Read the entire source row as a dict: col_index -> value
        src_row_vals: dict[int, object] = {}
        for cell in src_ws[src_row_idx]:
            src_row_vals[cell.column] = cell.value

        # Check if the source row is entirely blank — skip it
        if not any(v is not None for v in src_row_vals.values()):
            continue

        for t_col_offset, col_name in enumerate(template_columns):
            out_col = meta.header_col_start + t_col_offset
            src_col_idx = src_col_map.get(col_name)
            value = src_row_vals.get(src_col_idx) if src_col_idx is not None else None
            out_ws.cell(row=out_row, column=out_col).value = value

        out_row += 1

    buf = io.BytesIO()
    out_wb.save(buf)
    buf.seek(0)
    return buf.read(), warnings
